"""Parse real synthetic document bytes without external services or OCR."""

import io
import pytest
from euhnn.documents import iter_document, decode_text
from euhnn import ValidationError


def pdf_bytes(text="The blue wavelength is 0.46.", encrypted=False, blank=False):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    writer = PdfWriter()
    for n in range(2):
        page = writer.add_blank_page(width=612, height=792)
        if not blank:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
            )
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text} Page {n + 1}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("fixture-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_real_pdf_pages_and_citations(index):
    blob = pdf_bytes()
    pages = list(iter_document(blob, "guide.pdf"))
    assert len(pages) == 2 and [p.page for p in pages] == [1, 2]
    assert "Page 2" in pages[1].text
    index.ingest_bytes(blob, filename="guide.pdf")
    hits = index.search("blue wavelength")
    assert hits and all(h.page in (1, 2) and "page " in h.citation for h in hits)


@pytest.mark.parametrize("kwargs", [{"encrypted": True}, {"blank": True}])
def test_encrypted_and_textless_pdfs_fail_explicitly(kwargs):
    with pytest.raises(ValidationError):
        list(iter_document(pdf_bytes(**kwargs), "guide.pdf"))


def test_docx_preserves_paragraph_table_body_order():
    from docx import Document

    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "blue"
    table.cell(0, 1).text = "0.46"
    document.add_paragraph("After table")
    out = io.BytesIO()
    document.save(out)
    pages = list(iter_document(out.getvalue(), "guide.docx"))
    assert pages[0].text == "Before table\nblue\t0.46\nAfter table"
    assert pages[0].page is None


def test_html_does_not_execute_or_index_scripts_styles():
    data = b"<h1>Optical memory</h1><script>steal_secret()</script><style>hidden</style><p>Blue &amp; green.</p>"
    text = list(iter_document(data, "guide.html"))[0].text
    assert "Blue & green." in text and "steal_secret" not in text and "hidden" not in text


@pytest.mark.parametrize(
    "data",
    ["\u00f3ptica \u6f22\u5b57".encode("utf-8"), "\u00f3ptica \u6f22\u5b57".encode("utf-16"), b"\xef\xbb\xbfhello"],
)
def test_supported_unicode_encodings(data):
    assert decode_text(data)


@pytest.mark.parametrize(
    "filename,data",
    [
        ("bad.exe", b"data"),
        ("bad.pdf", b"invalid"),
        ("bad.docx", b"invalid"),
        ("bad.txt", b"\xff\xff"),
        ("blank.txt", b" \n"),
        ("nul.txt", b"abc\0def"),
    ],
)
def test_invalid_documents_are_not_silently_indexed(filename, data):
    with pytest.raises(ValidationError):
        list(iter_document(data, filename))
